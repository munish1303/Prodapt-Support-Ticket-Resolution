"""Experiment 9 analysis: candidate wrong-scenario signals, detector selection (tuning set only), test-set effects.

See experiments/e9_wrong_scenario.py for the question and data. Reads experiments/results/e9_features.jsonl.
"""

from __future__ import annotations

import json
from collections import Counter

import numpy as np

from evaluation.data import RESULTS_DIR, read_jsonl

TOP_K = 8


def auroc(pos: list[float], neg: list[float]) -> float | None:
    """P(score of a positive > score of a negative), ties count half."""
    if not pos or not neg:
        return None
    p, n = np.asarray(pos)[:, None], np.asarray(neg)[None, :]
    return float(((p > n).sum() + 0.5 * (p == n).sum()) / (p.size * n.size))


def signals(row: dict) -> dict[str, float]:
    """Runtime signals: computed from understanding + retrieval output only (never from scenario labels).
    Oriented so that HIGHER = MORE LIKELY WRONG SCENARIO."""
    top = row["top"][:TOP_K]
    kb = row["kb_candidates"]
    rel = sorted((s["relevance"] for s in top), reverse=True)
    ce = [s["ce"] for s in top]
    cats = Counter(s["category"] for s in top)
    products = set(row["products"] or [])
    fused_kb = kb[0] if kb else None
    best_ce_kb = max(kb[:5], key=lambda s: s["ce"]) if kb else None
    other_kb_ce = [s["ce"] for s in kb[1:5]]
    return {
        # retrieval strength (already in the confidence score, for reference)
        "neg_rel_top": -rel[0] if rel else 0.0,
        "neg_rel_mean5": -float(np.mean(rel[:5])) if rel else 0.0,
        "neg_intent_conf": -row["intent_confidence"],
        # ranker disagreement
        "ranker_disagreement": 1 - len(set(row["sem_top10"]) & set(row["lex_top10"])) / 10,
        "lex_rank_top1": float(top[0]["lex_rank"] or 999) if top else 999.0,
        # cross-encoder: complaint <-> sources
        "neg_ce_top1": -ce[0] if ce else 0.0,
        "neg_ce_mean": -float(np.mean(ce)) if ce else 0.0,
        "neg_ce_max": -max(ce) if ce else 0.0,
        "neg_ce_kb_fused": -fused_kb["ce"] if fused_kb else 0.0,
        # the cross-encoder prefers a different KB article than retrieval does
        "kb_disagree": float(bool(fused_kb and best_ce_kb and best_ce_kb["id"] != fused_kb["id"])),
        "kb_ce_margin": (max(other_kb_ce) - fused_kb["ce"]) if fused_kb and other_kb_ce else 0.0,
        # source consistency
        "category_split": 1 - cats.most_common(1)[0][1] / len(top) if top else 1.0,
        "intent_vs_sources": 1 - sum(s["category"] == row["intent"] for s in top) / len(top) if top else 1.0,
        "product_mismatch": (
            sum(1 for s in top if s["product"] and products and s["product"] not in products) / len(top) if top else 0.0
        ),
    }


def correct_share(row: dict) -> float:
    """Label: share of the production context (top 8) drawn from the complaint's true scenario."""
    top = row["top"][:TOP_K]
    return sum(s["scenario_id"] == row["scenario_id"] for s in top) / len(top) if top else 0.0


def wrong_context(row: dict) -> bool:
    return correct_share(row) < 0.5


def signal_table(rows: list[dict]) -> dict[str, float | None]:
    sig = [signals(r) for r in rows]
    lab = [wrong_context(r) for r in rows]
    out = {}
    for name in sig[0]:
        out[name] = auroc([s[name] for s, y in zip(sig, lab) if y], [s[name] for s, y in zip(sig, lab) if not y])
    return out


def explore() -> None:
    rows = read_jsonl(RESULTS_DIR / "e9_features.jsonl")
    by_set: dict[str, list[dict]] = {}
    for r in rows:
        by_set.setdefault(r["set"], []).append(r)
    for name in ("tune", "check"):
        rs = by_set.get(name, [])
        if not rs:
            continue
        n_wrong = sum(wrong_context(r) for r in rs)
        print(f"== {name}: n={len(rs)} wrong-context={n_wrong} ({n_wrong / len(rs):.0%})")
        for k, v in sorted(signal_table(rs).items(), key=lambda kv: -(kv[1] or 0)):
            print(f"   {k:20s} AUROC {v:.3f}" if v is not None else f"   {k:20s} n/a")


# Pre-registered selection rules (fixed before any test-set result was looked at):
CHEAP = [
    "neg_intent_conf",
    "intent_vs_sources",
    "category_split",
    "neg_rel_mean5",
    "neg_rel_top",
    "product_mismatch",
    "ranker_disagreement",
    "lex_rank_top1",
]  # no extra model at request time
CE = ["neg_ce_top1", "neg_ce_mean", "neg_ce_max", "neg_ce_kb_fused", "kb_disagree", "kb_ce_margin"]
CE_MIN_GAIN = 0.03  # use cross-encoder signals only if they add at least this much CV AUROC (they cost latency)
MAX_FALSE_FLAG = 0.10  # operating point: flag at most 10% of tuning cases whose context is correct
SEED = 0


def _matrix(rows: list[dict], names: list[str]) -> np.ndarray:
    return np.array([[signals(r)[n] for n in names] for r in rows], dtype=float)


def _fit(x: np.ndarray, y: np.ndarray):
    from sklearn.linear_model import LogisticRegression

    mu, sd = x.mean(0), x.std(0) + 1e-9
    model = LogisticRegression(C=1.0, max_iter=1000).fit((x - mu) / sd, y)
    return mu, sd, model


def _score(fitted, x: np.ndarray) -> np.ndarray:
    mu, sd, model = fitted
    return model.decision_function((x - mu) / sd)


def cv_auroc(rows: list[dict], names: list[str], folds: int = 5) -> float:
    rng = np.random.default_rng(SEED)
    idx = rng.permutation(len(rows))
    x, y = _matrix(rows, names), np.array([wrong_context(r) for r in rows])
    scores = np.zeros(len(rows))
    for f in range(folds):
        test = idx[f::folds]
        train = np.setdiff1d(idx, test)
        scores[test] = _score(_fit(x[train], y[train]), x[test])
    return auroc(list(scores[y]), list(scores[~y])) or 0.0


def select(tune: list[dict]) -> dict:
    """Greedy forward selection by 5-fold CV AUROC on the tuning set, within the cheap pool, then with the
    cross-encoder pool added; keep the cross-encoder only if it gains >= CE_MIN_GAIN."""

    def forward(pool: list[str]) -> tuple[list[str], float]:
        chosen: list[str] = []
        best = 0.0
        while True:
            trials = [(cv_auroc(tune, chosen + [c]), c) for c in pool if c not in chosen]
            if not trials:
                break
            score, cand = max(trials)
            if score <= best + 0.005:  # stop when a feature adds less than 0.005 AUROC
                break
            chosen, best = chosen + [cand], score
        return chosen, best

    cheap, cheap_auc = forward(CHEAP)
    full, full_auc = forward(CHEAP + CE)
    use_ce = full_auc >= cheap_auc + CE_MIN_GAIN
    names = full if use_ce else cheap
    x, y = _matrix(tune, names), np.array([wrong_context(r) for r in tune])
    fitted = _fit(x, y)
    s = _score(fitted, x)
    # threshold: the (1 - MAX_FALSE_FLAG) quantile of scores among correct-context tuning cases
    threshold = float(np.quantile(s[~y], 1 - MAX_FALSE_FLAG))
    mu, sd, model = fitted
    return {
        "cheap_features": cheap,
        "cheap_cv_auroc": round(cheap_auc, 4),
        "with_cross_encoder_features": full,
        "with_cross_encoder_cv_auroc": round(full_auc, 4),
        "uses_cross_encoder": use_ce,
        "features": names,
        "standardise_mean": [round(float(v), 6) for v in mu],
        "standardise_sd": [round(float(v), 6) for v in sd],
        "coef": [round(float(v), 6) for v in model.coef_[0]],
        "intercept": round(float(model.intercept_[0]), 6),
        "threshold": round(threshold, 6),
        "tune_flag_rate_wrong": round(float((s[y] >= threshold).mean()), 4),
        "tune_flag_rate_correct": round(float((s[~y] >= threshold).mean()), 4),
        "_fitted": fitted,
    }


def _rater_labels() -> dict[str, bool]:
    """case id -> AI rater's 'safe to use as-is' (EVALUATION.md §6.4)."""
    from openpyxl import load_workbook

    wb = load_workbook(RESULTS_DIR.parent.parent / "data" / "evaluation" / "ai_eval_sheet.xlsx")
    rat, sy = wb["Ratings"], wb["System"]
    return {sy.cell(row=r, column=2).value: rat.cell(row=r, column=9).value == "Yes" for r in range(3, 53)}


def _effect(flags: dict[str, bool], runs: dict[str, dict], bad) -> dict:
    """What happens if flagged RESOLVEs become REVIEW."""
    resolves = [i for i, r in runs.items() if r["decision"] == "RESOLVE" and i in flags]
    bad_res = [i for i in resolves if bad(i)]
    good_res = [i for i in resolves if not bad(i)]
    caught = [i for i in bad_res if flags[i]]
    lost = [i for i in good_res if flags[i]]
    n = len([i for i in runs if i in flags]) or 1
    after = len(resolves) - len(caught) - len(lost)
    return {
        "n": n,
        "resolves_before": len(resolves),
        "bad_resolves": len(bad_res),
        "bad_resolves_caught": len(caught),
        "good_resolves": len(good_res),
        "good_resolves_sent_to_review": len(lost),
        "resolve_rate_before": round(len(resolves) / n, 4),
        "resolve_rate_after": round(after / n, 4),
        "precision_of_resolve_before": round(len(good_res) / len(resolves), 4) if resolves else None,
        "precision_of_resolve_after": round((len(good_res) - len(lost)) / after, 4) if after else None,
        "caught_ids": sorted(caught),
        "lost_ids": sorted(lost),
    }


def analyze() -> None:
    rows = read_jsonl(RESULTS_DIR / "e9_features.jsonl")
    by_set: dict[str, list[dict]] = {}
    for r in rows:
        by_set.setdefault(r["set"], []).append(r)
    tune = by_set["tune"]
    sel = select(tune)
    fitted = sel.pop("_fitted")
    names = sel["features"]

    def flag_map(rs: list[dict]) -> tuple[dict[str, bool], dict[str, float]]:
        s = _score(fitted, _matrix(rs, names)) if rs else np.array([])
        return {r["id"]: bool(v >= sel["threshold"]) for r, v in zip(rs, s)}, {r["id"]: float(v) for r, v in zip(rs, s)}

    out: dict = {"selection": sel, "signal_auroc_tune": signal_table(tune)}
    check = by_set.get("check", [])
    if check:
        cs = _score(fitted, _matrix(check, names))
        y = np.array([wrong_context(r) for r in check])
        out["check_retrieval_level"] = {
            "n": len(check),
            "wrong_context": int(y.sum()),
            "detector_auroc": round(auroc(list(cs[y]), list(cs[~y])) or 0, 4),
            "flag_rate_wrong": round(float((cs[y] >= sel["threshold"]).mean()), 4),
            "flag_rate_correct": round(float((cs[~y] >= sel["threshold"]).mean()), 4),
            "signal_auroc": signal_table(check),
        }

    # draft-level tests
    gen_flags, gen_scores = flag_map(by_set.get("gen", []))
    nov_flags, _ = flag_map(by_set.get("novel", []))
    llm = json.load(open(RESULTS_DIR / "generation_llm.json", encoding="utf-8"))
    ext = json.load(open(RESULTS_DIR / "generation_extractive.json", encoding="utf-8"))
    llm_gen = {r["id"]: r for r in llm["rows"]}
    ext_gen = {r["id"]: r for r in ext["rows"]}
    rater = _rater_labels()
    out["test_generation_llm"] = _effect(gen_flags, llm_gen, lambda i: llm_gen[i]["reference_step_recall"] < 0.5)
    out["test_generation_extractive"] = _effect(gen_flags, ext_gen, lambda i: ext_gen[i]["reference_step_recall"] < 0.5)
    rated = {i: llm_gen[i] for i in rater}
    out["test_ai_rated_drafts"] = _effect(gen_flags, rated, lambda i: not rater[i])
    for name, run in (("llm", llm), ("extractive", ext)):
        nov = {r["id"]: r for r in run["novel_rows"]}
        out[f"test_novel_{name}"] = _effect(nov_flags, nov, lambda i: True)  # every RESOLVE of an unseen class is wrong
    ex = by_set.get("example", [])
    if ex:
        f, s = flag_map(ex)
        out["example_6_3"] = {"flagged": f[ex[0]["id"]], "score": round(s[ex[0]["id"]], 4)}

    # Ranking quality on the test sets (threshold-free): detector score of bad vs good RESOLVEs.
    def resolve_auroc(run: dict[str, dict], bad) -> float | None:
        res = [i for i, r in run.items() if r["decision"] == "RESOLVE" and i in gen_scores]
        a = auroc([gen_scores[i] for i in res if bad(i)], [gen_scores[i] for i in res if not bad(i)])
        return round(a, 4) if a is not None else None

    out["test_auroc_bad_vs_good_resolves"] = {
        "llm": resolve_auroc(llm_gen, lambda i: llm_gen[i]["reference_step_recall"] < 0.5),
        "extractive": resolve_auroc(ext_gen, lambda i: ext_gen[i]["reference_step_recall"] < 0.5),
        "ai_rated": resolve_auroc(rated, lambda i: not rater[i]),
    }
    out["gen_flag_rate"] = round(sum(gen_flags.values()) / len(gen_flags), 4) if gen_flags else None
    out["novel_flag_rate"] = round(sum(nov_flags.values()) / len(nov_flags), 4) if nov_flags else None
    (RESULTS_DIR / "e9_wrong_scenario.json").write_text(json.dumps(out, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in out.items() if not k.startswith("signal_auroc")}, indent=2))
    print("saved", RESULTS_DIR / "e9_wrong_scenario.json")


if __name__ == "__main__":
    explore()
